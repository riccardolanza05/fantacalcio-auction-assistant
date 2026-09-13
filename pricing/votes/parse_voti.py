"""
Parser v2 dei voti storici Fantacalcio.

Struttura reale del file (verificata):
    r1-r4  intestazione/disclaimer
    r5     nome squadra           <- blocco squadra
    r6     header: Cod. Ruolo Nome Voto Gf Gs Rp Rs Rf Au Amm Esp Ass [Gdv Gdp]
    r7..   righe giocatore
    ...    e si ripete per ogni squadra

Fogli: 'Fantacalcio', 'Statistico' (dal 2016-17 circa), 'Italia'.
I voti differiscono tra fogli: vengono tenuti separati nella colonna 'fonte'.

'Cod.' e' lo stesso identificativo usato nei file Quotazioni/Statistiche,
quindi il join e' diretto su quella chiave.

USO:
    python parse_voti_v2.py

Output:
    voti_panel.csv         una riga per (stagione, giornata, fonte, Id)
    voti_panel_report.txt  diagnostica
"""

import csv
import glob
import os
import re
import sys

try:
    import openpyxl
except ImportError:
    sys.exit("Serve openpyxl:  pip install openpyxl")

INDIR = "voti_raw"
OUT_CSV = "voti_panel.csv"
OUT_REPORT = "voti_panel_report.txt"

# Righe di intestazione da ignorare quando cerchiamo il nome squadra
BOILERPLATE = ("voti ", "solo su", "questo file", "e' da consider",
               "è da consider")

COLONNE_BASE = ["Cod.", "Ruolo", "Nome", "Voto", "Gf", "Gs", "Rp", "Rs",
                "Rf", "Au", "Amm", "Esp", "Ass"]

CAMPI_OUT = ["stagione", "giornata", "fonte", "Id", "Ruolo", "Nome", "Squadra",
             "Voto", "voto_ufficio", "senza_voto",
             "Gf", "Gs", "Rp", "Rs", "Rf", "Au", "Amm", "Esp", "Ass"]

RE_NUM = re.compile(r"-?\d+(?:[.,]\d+)?")


def parse_voto(v):
    """Ritorna (voto_float|None, ha_asterisco, senza_voto).

    Il voto puo' essere: 6.5 (numero), '6*' (voto d'ufficio: pochi minuti),
    's.v.'/'sv'/'-' (senza voto), oppure vuoto (non convocato/non entrato).
    """
    if v is None:
        return None, False, True
    if isinstance(v, (int, float)):
        return float(v), False, False

    s = str(v).strip()
    if not s:
        return None, False, True

    asterisco = "*" in s
    basso = s.lower().replace(" ", "")
    if basso in ("s.v.", "sv", "-", "n.d.", "nd"):
        return None, asterisco, True

    m = RE_NUM.search(s)
    if not m:
        return None, asterisco, True
    return float(m.group(0).replace(",", ".")), asterisco, False


def num(v):
    """Converte a int, tollerando None e stringhe."""
    if v is None or v == "":
        return 0
    if isinstance(v, (int, float)):
        return int(v)
    m = RE_NUM.search(str(v))
    return int(float(m.group(0).replace(",", "."))) if m else 0


def parse_foglio(ws, stagione, giornata, fonte, righe_out, problemi):
    """Scorre il foglio a stati: cerca blocchi squadra -> header -> giocatori."""
    righe = list(ws.iter_rows(values_only=True))
    squadra_corrente = None
    indici = None          # mappa nome_colonna -> indice
    n_giocatori = 0

    for i, riga in enumerate(righe):
        if not riga:
            continue
        c0 = riga[0]

        # --- riga di header? -> definisce le colonne e la squadra ---
        if isinstance(c0, str) and c0.strip().lower().startswith("cod"):
            intestazioni = [str(c).strip() if c is not None else "" for c in riga]
            indici = {h: j for j, h in enumerate(intestazioni) if h}

            # Il nome squadra e' nella riga precedente non vuota,
            # escludendo il boilerplate iniziale.
            squadra_corrente = None
            for k in range(i - 1, -1, -1):
                prec = righe[k]
                if not prec or prec[0] is None:
                    continue
                testo = str(prec[0]).strip()
                if any(testo.lower().startswith(b) for b in BOILERPLATE):
                    break
                # una riga squadra ha solo la prima cella popolata
                if all(c is None for c in prec[1:]):
                    squadra_corrente = testo
                break
            continue

        # --- riga giocatore? (prima cella numerica) ---
        if indici is None:
            continue
        if not isinstance(c0, (int, float)):
            continue

        def get(nome):
            j = indici.get(nome)
            return riga[j] if (j is not None and j < len(riga)) else None

        voto, asterisco, sv = parse_voto(get("Voto"))

        righe_out.append({
            "stagione": stagione,
            "giornata": giornata,
            "fonte": fonte,
            "Id": int(c0),
            "Ruolo": (str(get("Ruolo")).strip() if get("Ruolo") else ""),
            "Nome": (str(get("Nome")).strip() if get("Nome") else ""),
            "Squadra": squadra_corrente or "",
            "Voto": voto if voto is not None else "",
            "voto_ufficio": int(asterisco),
            "senza_voto": int(sv),
            "Gf": num(get("Gf")), "Gs": num(get("Gs")),
            "Rp": num(get("Rp")), "Rs": num(get("Rs")), "Rf": num(get("Rf")),
            "Au": num(get("Au")), "Amm": num(get("Amm")),
            "Esp": num(get("Esp")), "Ass": num(get("Ass")),
        })
        n_giocatori += 1

    if n_giocatori == 0:
        problemi.append(f"{stagione} g{giornata:02d} [{fonte}]: 0 giocatori estratti")
    return n_giocatori


def main():
    if not os.path.isdir(INDIR):
        sys.exit(f"Cartella '{INDIR}' assente. Lancia prima riordina_download.py")

    files = sorted(glob.glob(os.path.join(INDIR, "*", "g*.xlsx")))
    if not files:
        sys.exit(f"Nessun .xlsx in {INDIR}/")

    print(f"Trovati {len(files)} file. Elaborazione...")

    righe_out, problemi = [], []
    fogli_visti = {}
    senza_squadra = 0

    for n, path in enumerate(files, 1):
        stagione = os.path.basename(os.path.dirname(path))
        m = re.search(r"g(\d+)\.xlsx$", path)
        giornata = int(m.group(1)) if m else -1

        try:
            wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        except Exception as e:
            problemi.append(f"{stagione} g{giornata:02d}: illeggibile ({e})")
            continue

        for nome_foglio in wb.sheetnames:
            fogli_visti[nome_foglio] = fogli_visti.get(nome_foglio, 0) + 1
            parse_foglio(wb[nome_foglio], stagione, giornata,
                         nome_foglio, righe_out, problemi)
        wb.close()

        if n % 50 == 0:
            print(f"  {n}/{len(files)}  ({len(righe_out)} righe finora)")

    if not righe_out:
        sys.exit("Nessuna riga estratta. La struttura non corrisponde: "
                 "rilancia ispeziona_voti.py e manda l'output.")

    # deduplica su (stagione, giornata, fonte, Id)
    visti = set()
    finali = []
    dup = 0
    for r in righe_out:
        k = (r["stagione"], r["giornata"], r["fonte"], r["Id"])
        if k in visti:
            dup += 1
            continue
        visti.add(k)
        if not r["Squadra"]:
            senza_squadra += 1
        finali.append(r)

    with open(OUT_CSV, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=CAMPI_OUT)
        w.writeheader()
        w.writerows(finali)

    # ---------------- report ----------------
    stagioni = {}
    for r in finali:
        s = stagioni.setdefault(r["stagione"], {"righe": 0, "giornate": set(),
                                                "id": set(), "fonti": set()})
        s["righe"] += 1
        s["giornate"].add(r["giornata"])
        s["id"].add(r["Id"])
        s["fonti"].add(r["fonte"])

    L = [
        "REPORT PARSING VOTI (v2)",
        "=" * 66,
        f"File elaborati   : {len(files)}",
        f"Righe totali     : {len(finali)}   (duplicati scartati: {dup})",
        f"Giocatori unici  : {len({r['Id'] for r in finali})}",
        f"Righe senza squadra: {senza_squadra}",
        "",
        "FOGLI TROVATI (n. file in cui compaiono):",
    ]
    for k, v in sorted(fogli_visti.items()):
        L.append(f"  {k}: {v}")

    L += ["", "COPERTURA PER STAGIONE:",
          f"  {'stagione':10s} {'giornate':>9s} {'righe':>8s} {'giocatori':>10s}  fonti"]
    for st in sorted(stagioni):
        s = stagioni[st]
        L.append(f"  {st:10s} {len(s['giornate']):>9d} {s['righe']:>8d} "
                 f"{len(s['id']):>10d}  {','.join(sorted(s['fonti']))}")

    # qualita' dei voti
    tot = len(finali)
    n_sv = sum(r["senza_voto"] for r in finali)
    n_ast = sum(r["voto_ufficio"] for r in finali)
    L += ["", "QUALITA' VOTI:",
          f"  con voto numerico : {tot - n_sv} ({100*(tot-n_sv)/tot:.1f}%)",
          f"  senza voto        : {n_sv} ({100*n_sv/tot:.1f}%)",
          f"  voto d'ufficio (*): {n_ast} ({100*n_ast/tot:.1f}%)"]

    if problemi:
        L += ["", f"PROBLEMI ({len(problemi)}):"] + [f"  {p}" for p in problemi[:40]]
        if len(problemi) > 40:
            L.append(f"  ... e altri {len(problemi) - 40}")

    testo = "\n".join(L)
    with open(OUT_REPORT, "w", encoding="utf-8") as f:
        f.write(testo)

    print("\n" + testo)
    print(f"\n-> {OUT_CSV}\n-> {OUT_REPORT}")
    print("\nProssimo passo: python valida_voti.py")


if __name__ == "__main__":
    main()