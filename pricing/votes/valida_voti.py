"""
Valida voti_panel.csv confrontandolo con le statistiche ufficiali.

Due verifiche:

1) PRESENZE. Ricalcola le presenze di ogni giocatore dai voti giornata per
   giornata e le confronta con la colonna Pv dei file Statistiche. Prova due
   definizioni per capire quale usa fantacalcio.it:
       A) ogni voto numerico conta        (asterisco incluso)
       B) il voto d'ufficio non conta     (asterisco escluso)
   Quella con l'errore piu' basso e' la definizione corretta.

2) ARRIVI DI GENNAIO. Individua la prima giornata in cui ogni giocatore
   compare: da qui si ricava quante partite aveva realmente a disposizione,
   che era il problema del censoring rimasto aperto.

USO:
    python valida_voti.py [cartella_statistiche]

Se non passi la cartella, cerca i file Statistiche_*.xlsx nella cartella
corrente.
"""

import csv
import glob
import os
import re
import sys
from collections import defaultdict

try:
    import openpyxl
except ImportError:
    sys.exit("Serve openpyxl:  pip install openpyxl")

CSV_VOTI = "voti_panel.csv"
FONTE = "Fantacalcio"      # foglio di riferimento per le leghe standard
OUT = "validazione_voti.txt"


def cartella_stats():
    return sys.argv[1] if len(sys.argv) > 1 else "."


def carica_voti():
    if not os.path.exists(CSV_VOTI):
        sys.exit(f"{CSV_VOTI} assente. Lancia prima parse_voti_v2.py")

    # per (stagione, Id): presenze secondo le due definizioni, prima/ultima giornata
    dati = defaultdict(lambda: {"A": 0, "B": 0, "prima": 99, "ultima": 0,
                                "nome": "", "giornate": set()})
    fonti = set()

    with open(CSV_VOTI, encoding="utf-8") as f:
        for r in csv.DictReader(f):
            fonti.add(r["fonte"])
            if r["fonte"] != FONTE:
                continue

            k = (r["stagione"], int(r["Id"]))
            d = dati[k]
            d["nome"] = r["Nome"]
            g = int(r["giornata"])

            ha_voto = r["Voto"] not in ("", None) and r["senza_voto"] != "1"
            asterisco = r["voto_ufficio"] == "1"

            if ha_voto:
                d["A"] += 1
                if not asterisco:
                    d["B"] += 1
            # la comparsa in lista conta per capire da quando e' in rosa
            d["giornate"].add(g)
            d["prima"] = min(d["prima"], g)
            d["ultima"] = max(d["ultima"], g)

    if not dati:
        sys.exit(f"Nessuna riga con fonte '{FONTE}'. Fonti trovate: {sorted(fonti)}")
    return dati, fonti


def carica_pv(dirstats):
    pattern = os.path.join(dirstats, "Statistiche_Fantacalcio_Stagione_*.xlsx")
    files = sorted(glob.glob(pattern))
    if not files:
        sys.exit(f"Nessun file Statistiche trovato in: {os.path.abspath(dirstats)}\n"
                 f"Passa la cartella: python valida_voti.py \"C:/percorso\"")

    pv = {}
    for path in files:
        m = re.search(r"(\d{4})_(\d{2})", os.path.basename(path))
        if not m:
            continue
        stagione = f"{m.group(1)}-{m.group(2)}"
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        ws = wb["Tutti"] if "Tutti" in wb.sheetnames else wb[wb.sheetnames[0]]

        righe = ws.iter_rows(values_only=True)
        header = None
        for riga in righe:
            if riga and riga[0] == "Id":
                header = {str(c).strip(): j for j, c in enumerate(riga) if c}
                continue
            if header is None or not riga or not isinstance(riga[0], (int, float)):
                continue
            i_pv = header.get("Pv")
            if i_pv is None:
                break
            pv[(stagione, int(riga[0]))] = int(riga[i_pv] or 0)
        wb.close()
    return pv


def main():
    dati, fonti = carica_voti()
    pv = carica_pv(cartella_stats())

    L = ["VALIDAZIONE VOTI", "=" * 66,
         f"fonte usata: {FONTE}   (fonti disponibili: {', '.join(sorted(fonti))})",
         f"coppie (stagione, giocatore) dai voti: {len(dati)}",
         f"coppie con Pv ufficiale disponibile  : {len(pv)}", ""]

    # ---------- 1) presenze ----------
    comuni = [k for k in dati if k in pv]
    if not comuni:
        L.append("ATTENZIONE: nessuna corrispondenza tra voti e statistiche.")
        L.append("Gli identificativi non combaciano: verifica le stagioni scaricate.")
    else:
        errA = errB = 0
        exactA = exactB = 0
        for k in comuni:
            a, b, p = dati[k]["A"], dati[k]["B"], pv[k]
            errA += abs(a - p)
            errB += abs(b - p)
            exactA += (a == p)
            exactB += (b == p)
        n = len(comuni)
        L += [f"1) PRESENZE — confronto su {n} coppie",
              f"   def. A (asterisco conta)    : MAE {errA/n:.3f}  esatte {100*exactA/n:.1f}%",
              f"   def. B (asterisco NON conta): MAE {errB/n:.3f}  esatte {100*exactB/n:.1f}%",
              ""]
        migliore = "A" if errA <= errB else "B"
        L.append(f"   -> definizione corretta: {migliore}")
        if min(errA, errB) / n < 0.05:
            L.append("   -> ricostruzione praticamente perfetta: i dati sono affidabili.")
        elif min(errA, errB) / n < 0.5:
            L.append("   -> scarto piccolo: probabili giornate mancanti o rinvii.")
        else:
            L.append("   -> scarto ALTO: qualcosa non torna, da investigare.")
        L.append("")

        # esempi di discrepanza
        peggiori = sorted(comuni,
                          key=lambda k: -abs(dati[k][migliore] - pv[k]))[:10]
        L.append("   maggiori discrepanze:")
        for k in peggiori:
            st, idg = k
            L.append(f"     {st} {dati[k]['nome'][:22]:22s} "
                     f"calcolate {dati[k][migliore]:2d} vs Pv {pv[k]:2d}")
        L.append("")

    # ---------- 2) arrivi di gennaio ----------
    per_stagione = defaultdict(lambda: {"tot": 0, "tardivi": 0, "usciti": 0})
    esempi = []
    for (st, idg), d in dati.items():
        s = per_stagione[st]
        s["tot"] += 1
        if d["prima"] >= 19:
            s["tardivi"] += 1
            if len(esempi) < 12:
                esempi.append((st, d["nome"], d["prima"], d["ultima"]))
        if d["ultima"] <= 19:
            s["usciti"] += 1

    L += ["2) DISPONIBILITA' REALE (censoring)",
          "   giocatori la cui prima comparsa e' dalla giornata 19 in poi",
          "   = arrivati nel mercato invernale",
          "",
          f"   {'stagione':10s} {'giocatori':>10s} {'tardivi':>8s} {'%':>6s} {'usciti':>7s}"]
    for st in sorted(per_stagione):
        s = per_stagione[st]
        pct = 100 * s["tardivi"] / s["tot"] if s["tot"] else 0
        L.append(f"   {st:10s} {s['tot']:>10d} {s['tardivi']:>8d} {pct:>5.1f}% {s['usciti']:>7d}")

    if esempi:
        L += ["", "   esempi (prima -> ultima giornata):"]
        for st, nome, p, u in esempi:
            L.append(f"     {st} {nome[:22]:22s} g{p:02d} -> g{u:02d}")

    testo = "\n".join(L)
    with open(OUT, "w", encoding="utf-8") as f:
        f.write(testo)
    print(testo)
    print(f"\n-> {OUT}")


if __name__ == "__main__":
    main()