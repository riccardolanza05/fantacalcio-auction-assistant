"""
Parser per i file "Rose_lega-*.xlsx" esportati da FantaAsta Live / fantacalcio.it.
Struttura: blocchi di 2 squadre affiancate, ripetuti verticalmente.
Ogni blocco ha: riga nome-squadra, riga header (Ruolo/Calciatore/Squadra/Costo),
poi N righe giocatore, poi una riga "Crediti Residui: X".

Ritorna una lista di dict: {squadra_fantacalcio, ruolo, calciatore, squadra_reale, costo}
"""
import openpyxl

VALID_ROLES = {"P", "D", "C", "A"}


def parse_rose_lega(path: str, sheet_name: str = "TutteLeRose") -> list[dict]:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet_name]

    records = []
    # due blocchi colonna: (1..4) e (6..9)
    col_blocks = [(1, 2, 3, 4), (6, 7, 8, 9)]

    current_team = {0: None, 1: None}
    for r in range(1, ws.max_row + 1):
        row_vals = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
        for block_idx, (c_ruolo, c_nome, c_squadra, c_costo) in enumerate(col_blocks):
            ruolo = ws.cell(row=r, column=c_ruolo).value
            nome = ws.cell(row=r, column=c_nome).value
            squadra = ws.cell(row=r, column=c_squadra).value
            costo = ws.cell(row=r, column=c_costo).value

            # riga di intestazione squadra fantacalcio: es. ('NapoLanza', None, None, None)
            if ruolo and not costo and ruolo not in ("Ruolo",) and nome is None:
                current_team[block_idx] = str(ruolo).strip()
                continue

            if ruolo in VALID_ROLES and nome and squadra and costo is not None:
                try:
                    costo_val = int(str(costo).strip())
                except ValueError:
                    continue
                records.append({
                    "squadra_fantacalcio": current_team[block_idx],
                    "ruolo": str(ruolo).strip(),
                    "calciatore": str(nome).strip(),
                    "squadra_reale": str(squadra).strip(),
                    "costo": costo_val,
                })
    return records


if __name__ == "__main__":
    import sys

    if len(sys.argv) != 2:
        print("Uso: python parse_rose_lega.py <percorso a Rose_lega-*.xlsx>")
        raise SystemExit(1)
    recs = parse_rose_lega(sys.argv[1])
    print(f"Totale giocatori estratti: {len(recs)}")
    teams = sorted(set(r["squadra_fantacalcio"] for r in recs))
    print(f"Squadre fantacalcio trovate ({len(teams)}): {teams}")
    from collections import Counter
    print("Distribuzione ruoli:", Counter(r["ruolo"] for r in recs))
    print("Esempio:", recs[:3])
