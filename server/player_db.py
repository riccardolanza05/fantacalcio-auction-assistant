"""
Indice dei giocatori 2026-27 con fair value / stime del modello, per lookup
rapido durante l'asta (nome + squadra reale -> tutte le stats rilevanti).
"""
import openpyxl
from common import normalize_name, team_to_abbr, build_match_key


class PlayerDB:
    def __init__(self, fair_values_xlsx: str):
        self.by_key = {}
        self.by_name = {}
        self.by_id = {}
        self.by_surname = {}
        self._load(fair_values_xlsx)
        self._build_name_index()

    def _load(self, path: str):
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb[wb.sheetnames[0]]
        header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        idx = {h: i for i, h in enumerate(header)}

        def g(row, col):
            return row[idx[col]] if col in idx else None

        for row in ws.iter_rows(min_row=2, values_only=True):
            nome = g(row, "Nome")
            if nome is None:
                continue
            squadra = g(row, "Squadra") or ""
            abbr = team_to_abbr(squadra)
            record = {
                "id": g(row, "Id"),
                "nome": nome,
                "squadra": squadra,
                "ruolo": g(row, "Ruolo"),
                "fascia": g(row, "fascia"),
                "pct_titolarita": g(row, "pct_titolarita"),
                "fvm_mercato": g(row, "FVM"),
                "qti": g(row, "QtI"),
                "produzione_attesa": g(row, "produzione_attesa"),
                "p10": g(row, "p10"),
                "p90": g(row, "p90"),
                "prezzo_base": g(row, "prezzo_base"),
                "correzione_statistica": g(row, "correzione"),
                "prezzo_modello": g(row, "prezzo_modello"),
                "prezzo_mercato": g(row, "prezzo_mercato"),
                "scarto_pct": g(row, "scarto_pct"),
                "verdetto": g(row, "verdetto"),
                "rif_verdetto": g(row, "rif_verdetto"),
                "gol_2025_26": g(row, "gol_2025_26"),
                "gol_subiti_2025_26": g(row, "gol_subiti_2025_26"),
                "clean_sheet_2025_26": g(row, "clean_sheet_2025_26"),
                "rigori_parati_2025_26": g(row, "rigori_parati_2025_26"),
                "assist_2025_26": g(row, "assist_2025_26"),
                "presenze_2025_26": g(row, "presenze_2025_26"),
                "fm_media_2025_26": g(row, "fm_media_2025_26"),
                "nota_infortunio": g(row, "nota_infortunio"),
                "infortunato": bool(g(row, "infortunato")),
            }
            key = build_match_key(nome, abbr)
            self.by_key[key] = record
            self.by_name.setdefault(normalize_name(nome), []).append(record)
            if record["id"] is not None:
                self.by_id[int(record["id"])] = record

    def lookup_by_id(self, player_id) -> dict | None:
        """Match ESATTO via l'id fantacalcio.it (estratto dall'URL della card
        immagine nella pagina d'asta). Da preferire sempre al match per nome:
        la card mostra il nome esteso ("Lautaro Martinez") mentre i file usano
        la forma abbreviata ("Martinez L."), quindi il match testuale
        fallirebbe proprio sui giocatori piu' importanti."""
        try:
            return self.by_id.get(int(player_id))
        except (TypeError, ValueError):
            return None

    # ------------------------------------------------------------------
    # Matching per nome: la pagina d'asta mostra il nome ESTESO
    # ("Moise Kean", "Lautaro Martinez") mentre i file fantacalcio.it usano
    # la forma abbreviata ("Kean", "Martinez L."). Un confronto diretto
    # fallirebbe su quasi tutti i giocatori, quindi confrontiamo cognome +
    # iniziale del nome proprio.
    # ------------------------------------------------------------------
    @staticmethod
    def _split_db_name(nome_db: str):
        """'Martinez L.' -> ('martinez', 'l');  'Kean' -> ('kean', None)
        'Bella-Kotchap' -> ('bella kotchap', None)"""
        n = normalize_name(nome_db)          # minuscolo, senza accenti/punti
        parts = n.split()
        if len(parts) >= 2 and len(parts[-1]) <= 2:
            # ultimo token corto = iniziale del nome proprio (L., Jo., ...)
            return " ".join(parts[:-1]), parts[-1]
        return n, None

    @staticmethod
    def _split_full_name(nome_esteso: str):
        """'Lautaro Martinez' -> ('martinez', 'lautaro')
        'Moise Kean' -> ('kean', 'moise');  'Kean' -> ('kean', None)"""
        n = normalize_name(nome_esteso)
        parts = n.split()
        if len(parts) == 1:
            return parts[0], None
        return parts[-1], parts[0]

    def _build_name_index(self):
        """Indice cognome -> lista di record, costruito una volta sola."""
        self.by_surname = {}
        for rec in self.by_key.values():
            surname, _ = self._split_db_name(rec["nome"])
            self.by_surname.setdefault(surname, []).append(rec)

    def lookup(self, nome: str, squadra: str = "") -> dict | None:
        """Cerca per nome (abbreviato O esteso) + squadra."""
        if not nome:
            return None
        abbr = team_to_abbr(squadra) if squadra else ""

        # 1) match esatto nome+squadra (funziona se il nome e' gia' abbreviato,
        #    es. letto dalla lista laterale)
        if abbr:
            rec = self.by_key.get(build_match_key(nome, abbr))
            if rec:
                return rec

        # 2) match esatto solo nome, se univoco
        cands = self.by_name.get(normalize_name(nome), [])
        if len(cands) == 1:
            return cands[0]
        if len(cands) > 1 and abbr:
            same_team = [c for c in cands if team_to_abbr(c["squadra"]) == abbr]
            if len(same_team) == 1:
                return same_team[0]

        # 3) match cognome + iniziale (nome esteso della card d'asta)
        surname, firstname = self._split_full_name(nome)
        cands = list(self.by_surname.get(surname, []))
        if abbr:
            same_team = [c for c in cands if team_to_abbr(c["squadra"]) == abbr]
            if same_team:
                cands = same_team
        if len(cands) == 1:
            return cands[0]
        if len(cands) > 1 and firstname:
            # piu' omonimi di cognome: disambigua con l'iniziale del nome
            # ('Martinez L.' vs 'Martinez Jo.' contro 'Lautaro Martinez')
            exact = []
            for c in cands:
                _, initial = self._split_db_name(c["nome"])
                if initial and firstname.startswith(initial):
                    exact.append(c)
            if len(exact) == 1:
                return exact[0]
        return None

if __name__ == "__main__":
    import os
    HERE = os.path.dirname(os.path.abspath(__file__))
    db = PlayerDB(os.path.join(HERE, "giocatori_2026_27.xlsx"))
    print(f"Giocatori indicizzati: {len(db.by_key)} (per id: {len(db.by_id)})\n")
    for nome, squadra in [("Moise Kean", "Fiorentina"), ("Lautaro Martinez", "Inter"),
                           ("Josep Martinez", "Inter"), ("Kean", "Fiorentina")]:
        r = db.lookup(nome, squadra)
        if r:
            print(f"{nome:20s} -> {r['nome']:14s} prezzo {r['prezzo_modello']:6.1f} "
                  f"fascia {r['fascia']:8s} gol {r['gol_2025_26']} ass {r['assist_2025_26']} "
                  f"inf {r['infortunato']}")
    print("\nInfortunati nel database:")
    for r in list(db.by_id.values()):
        if r["infortunato"]:
            print(f"  {r['nome']:16s} {r['squadra']:12s} "
                  f"{r['nota_infortunio']}")
